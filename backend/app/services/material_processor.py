from __future__ import annotations

import asyncio
import hashlib
import logging
import re
from datetime import datetime, timezone
from pathlib import Path

import pdfplumber
import mammoth
from pptx import Presentation
from PIL import Image

try:
    import pytesseract
except ImportError:
    pytesseract = None  # type: ignore[assignment]

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import settings
from app.database import async_session_factory
from app.models import Chunk, Material, MaterialStatus, Topic, Difficulty

logger = logging.getLogger(__name__)

CHUNK_SIZE = 400
CHUNK_OVERLAP = 50


async def process_material(material_id: str) -> None:
    """Orchestrate full extraction pipeline for a material record."""
    async with async_session_factory() as db:
        result = await db.execute(select(Material).where(Material.id == material_id))
        material = result.scalar_one_or_none()
        if material is None:
            logger.error("Material %s not found", material_id)
            return

        try:
            material.status = MaterialStatus.processing
            await db.commit()

            file_path = material.storage_path
            file_type = material.file_type.lower()

            if file_type == "pdf":
                raw_text = extract_text_from_pdf(file_path)
            elif file_type == "docx":
                raw_text = extract_text_from_docx(file_path)
            elif file_type == "pptx":
                raw_text = extract_text_from_pptx(file_path)
            elif file_type in ("png", "jpg", "jpeg", "tiff", "bmp"):
                raw_text = extract_text_from_image(file_path)
            else:
                raise ValueError(f"Unsupported file type: {file_type}")

            cleaned = clean_text(raw_text)
            material.raw_text = cleaned

            chunks_text = chunk_text(cleaned, CHUNK_SIZE, CHUNK_OVERLAP)

            topics_list = _split_into_topics(cleaned, material.source_filename)

            canonical = build_canonical_json(material, topics_list, chunks_text)
            material.extracted_json = canonical

            topic_objs: list[Topic] = []
            for i, t in enumerate(topics_list):
                diff_raw = t.get("difficulty", "beginner")
                difficulty = Difficulty(diff_raw) if isinstance(diff_raw, str) else diff_raw
                topic = Topic(
                    material_id=material.id,
                    subject_id=material.subject_id,
                    title=t["title"],
                    subtopics=t.get("subtopics", []),
                    difficulty=difficulty,
                    order_index=i,
                )
                db.add(topic)
                topic_objs.append(topic)
            await db.flush()

            embedding_vectors: list[list[float]] = []
            if settings.GEMINI_API_KEY:
                embedding_vectors = await generate_embeddings(chunks_text)

            for i, chunk in enumerate(chunks_text):
                topic_idx = min(i // max(1, len(chunks_text) // max(1, len(topic_objs))), len(topic_objs) - 1)
                token_count = len(chunk.split())
                keywords = _extract_simple_keywords(chunk)
                chunk_obj = Chunk(
                    topic_id=topic_objs[topic_idx].id,
                    material_id=material.id,
                    text=chunk,
                    page_number=(i // 3) + 1,
                    keywords=keywords,
                    token_count=token_count,
                )
                if embedding_vectors and i < len(embedding_vectors):
                    chunk_obj.embedding = embedding_vectors[i]
                db.add(chunk_obj)

            material.status = MaterialStatus.processed
            material.processed_at = datetime.now(timezone.utc)
            await db.commit()
            logger.info("Material %s processed successfully", material_id)

        except Exception as exc:
            logger.exception("Failed to process material %s", material_id)
            material.status = MaterialStatus.failed
            await db.commit()


def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from a PDF using pdfplumber."""
    pages: list[str] = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            tables = page.extract_tables() or []
            for table in tables:
                for row in table:
                    cells = [str(cell) if cell else "" for cell in row]
                    text += "\n" + " | ".join(cells)
            pages.append(text)
    return "\n\n".join(pages)


def extract_text_from_docx(file_path: str) -> str:
    """Extract text from a DOCX using mammoth."""
    with open(file_path, "rb") as f:
        result = mammoth.convert_to_markdown(f)
    return result.value


def extract_text_from_pptx(file_path: str) -> str:
    """Extract text from a PPTX file."""
    prs = Presentation(file_path)
    texts: list[str] = []
    for slide in prs.slides:
        slide_text: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    slide_text.append(para.text)
            if shape.has_table:
                for row in shape.table.rows:
                    cells = [cell.text for cell in row.cells]
                    slide_text.append(" | ".join(cells))
        texts.append("\n".join(slide_text))
    return "\n\n".join(texts)


def extract_text_from_image(file_path: str) -> str:
    """Extract text from an image using Tesseract OCR."""
    if pytesseract is None:
        raise RuntimeError("pytesseract is not installed")
    img = Image.open(file_path)
    text = pytesseract.image_to_string(img)
    return text


def clean_text(raw_text: str) -> str:
    """Remove artifacts and normalize whitespace."""
    text = raw_text
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = re.sub(r"^\s+|\s+$", "", text, flags=re.MULTILINE)
    text = text.strip()
    return text


def chunk_text(text: str, chunk_size: int = CHUNK_SIZE, overlap: int = CHUNK_OVERLAP) -> list[str]:
    """Split text into overlapping chunks by word count."""
    words = text.split()
    if len(words) <= chunk_size:
        return [text] if text.strip() else []
    chunks: list[str] = []
    start = 0
    while start < len(words):
        end = start + chunk_size
        chunk = " ".join(words[start:end])
        chunks.append(chunk)
        start += chunk_size - overlap
    return chunks


async def generate_embeddings(chunks: list[str]) -> list[list[float]]:
    """Generate embeddings using Gemini API. Returns list of vectors."""
    try:
        import google.generativeai as genai

        genai.configure(api_key=settings.GEMINI_API_KEY)
        all_vectors: list[list[float]] = []
        batch_size = 100
        for i in range(0, len(chunks), batch_size):
            batch = chunks[i:i + batch_size]
            result = genai.embed_content(
                model="models/text-embedding-004",
                content=batch,
            )
            vectors = result["embedding"]  # type: ignore[index]
            all_vectors.extend(vectors)
        return all_vectors
    except Exception as exc:
        logger.warning("Embedding generation failed: %s", exc)
        return []


def tag_content(chunks: list[str], grade: int, subject: str) -> list[dict]:
    """Classify and extract keywords from chunks."""
    tagged: list[dict] = []
    for chunk in chunks:
        words = chunk.lower().split()
        freq: dict[str, int] = {}
        stop_words = {"the", "a", "an", "is", "are", "was", "were", "in", "on", "at", "to", "for", "of", "and", "or", "but", "with", "this", "that", "it"}
        for w in words:
            w_clean = re.sub(r"[^a-z0-9]", "", w)
            if len(w_clean) > 2 and w_clean not in stop_words:
                freq[w_clean] = freq.get(w_clean, 0) + 1
        top_keywords = sorted(freq.keys(), key=lambda k: freq[k], reverse=True)[:10]
        tagged.append({"text": chunk[:100], "keywords": top_keywords, "grade": grade, "subject": subject})
    return tagged


def _extract_simple_keywords(text: str, n: int = 10) -> list[str]:
    """Quick keyword extraction by frequency."""
    words = text.lower().split()
    stop_words = {"the", "a", "an", "is", "are", "was", "were", "in", "on", "at", "to", "for", "of", "and", "or", "but", "with", "this", "that", "it", "by", "from", "as"}
    freq: dict[str, int] = {}
    for w in words:
        w_clean = re.sub(r"[^a-z0-9]", "", w)
        if len(w_clean) > 2 and w_clean not in stop_words:
            freq[w_clean] = freq.get(w_clean, 0) + 1
    return sorted(freq.keys(), key=lambda k: freq[k], reverse=True)[:n]


def _split_into_topics(text: str, filename: str) -> list[dict]:
    """Split extracted text into topic dicts."""
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    if not paragraphs:
        return [{"title": Path(filename).stem.replace("_", " ").title(), "subtopics": [], "difficulty": "beginner"}]

    topics_list: list[dict] = []
    batch_size = max(1, len(paragraphs) // 3)
    for i in range(0, len(paragraphs), batch_size):
        batch = paragraphs[i: i + batch_size]
        first_line = batch[0][:120]
        topics_list.append({
            "title": first_line.split(".")[0].strip()[:200] or f"Topic {i // batch_size + 1}",
            "subtopics": [p[:80].split(".")[0].strip() for p in batch[:5]],
            "difficulty": "beginner",
        })
    return topics_list


def build_canonical_json(material: Material, topics: list[dict], chunks: list[str]) -> dict:
    """Assemble canonical JSON representation of processed material."""
    return {
        "material_id": str(material.id),
        "source_filename": material.source_filename,
        "file_type": material.file_type,
        "grade": material.grade,
        "subject_id": str(material.subject_id),
        "total_chunks": len(chunks),
        "topics": topics,
        "chunks_preview": [c[:200] for c in chunks[:5]],
        "checksum": material.checksum_sha256,
        "processed_at": datetime.now(timezone.utc).isoformat(),
    }
